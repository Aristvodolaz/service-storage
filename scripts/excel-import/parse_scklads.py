# -*- coding: utf-8 -*-
"""
Парсит выгрузку "1383 - все ячейки" и готовит справочник ячеек для
[SPOe_rc].[dbo].[x_Storage_Scklads] (WR_House = 1383).

Нужны только 3 поля таблицы: Name, SHK, WR_House (ID = IDENTITY).

    Name     <- колонка "Код"      (напр. 01-1-1)
    SHK      <- колонка "BarCode"  (напр. 2002913494)
    WR_House <- колонка "Склад #"  (1383)

Запуск:
    python parse_scklads.py "C:\\...\\1383 - все ячейки (1).xlsx" [out.json]
"""
import sys, os, json, collections

DEFAULT_XLSX = r"C:\Users\G15\Downloads\1383 - все ячейки (1).xlsx"
COL_CODE, COL_BARCODE, COL_SKLAD = 9, 12, 6


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    if not os.path.exists(xlsx):
        sys.exit("Файл не найден: " + xlsx)

    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = ws.iter_rows(values_only=True)
    next(rows, None)

    out = []
    stats = collections.Counter()
    names = set()
    dup_names, dup_barcodes = [], []
    barcodes = collections.Counter()

    for r in rows:
        if r is None or not any(c is not None for c in r):
            continue
        stats["total"] += 1
        code = r[COL_CODE]
        bc = r[COL_BARCODE]
        sk = r[COL_SKLAD]

        name = str(code).strip() if code is not None else None
        shk = str(bc).strip() if bc not in (None, "") else None
        house = str(int(sk)) if sk is not None else None

        if not name:
            stats["no_name"] += 1
            continue
        if name in names:
            dup_names.append(name)
            stats["dup_name"] += 1
        names.add(name)
        if shk is None:
            stats["no_barcode"] += 1
        else:
            barcodes[shk] += 1

        out.append({"name": name, "shk": shk, "house": house})

    for b, c in barcodes.items():
        if c > 1:
            dup_barcodes.append(b)

    here = os.path.dirname(os.path.abspath(__file__))
    out_name = sys.argv[2] if len(sys.argv) > 2 else "scklads_1383.json"
    out_path = out_name if os.path.isabs(out_name) else os.path.join(here, out_name)
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)

    print("Файл:            ", xlsx)
    print("Строк:           ", stats["total"])
    print("В справочник:     ", len(out))
    print("Без Кода (пропущены):", stats["no_name"])
    print("Без BarCode (SHK=NULL):", stats["no_barcode"])
    print("Дубли Кода:       ", stats["dup_name"], dup_names[:10])
    print("Дубли BarCode:    ", len(dup_barcodes), dup_barcodes[:10])
    print("->", out_path)


if __name__ == "__main__":
    main()
